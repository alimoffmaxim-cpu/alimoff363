"""Собирает шаблон отчёта «Сова» (xlsx) с демо-данными.

Все листы отчёта считаются формулами из листа «Данные» (день × направление),
поэтому при загрузке новых строк отчёт и графики обновляются сами.
Запуск: python3 build_xlsx.py out.xlsx
"""
import random
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NAVY, HEAD, BAND, LINE, INK, MUTED = '1F2A44', '2F3E60', 'F4F6FA', 'DFE3EB', '1F2430', '6B7280'
SPEND, LEADS, CPL = '2A78D6', '1BAF7A', 'EB6834'
DIR_COLORS = ['2A78D6', 'EB6834', '1BAF7A']
DIRS = ['Направление 1', 'Направление 2', 'Направление 3']
FONT = 'Arial'
D = "'Данные'"
SET = "'⚙ Настройки'"
COLS = ['Период', 'Расход', 'Показы', 'Клики', 'CTR', 'CPC', 'Лиды', 'CPL', 'CR в лид']
FMT = [None, '#,##0" ₽"', '#,##0', '#,##0', '0.00%', '#,##0.0" ₽"', '#,##0', '#,##0" ₽"', '0.0%']
DELTA_FMT = '"▲ "0%;"▼ "0%;0%'
GOOD = [0, 0, 1, 1, 1, -1, 1, -1, 1]  # +1 — хорошо, когда растёт; −1 — когда падает
N_MONTHS, N_WEEKS, N_DAYS = 12, 26, 60

thin = Side(style='thin', color=LINE)


def fill(c):
    return PatternFill('solid', start_color=c, end_color=c)


def font(**kw):
    kw.setdefault('name', FONT)
    kw.setdefault('color', INK)
    kw.setdefault('size', 10)
    return Font(**kw)


def col(i):
    from openpyxl.utils import get_column_letter
    return get_column_letter(i)


# ---------- формулы ----------

def sumifs(src, frm, to, dir_ref):
    crit = f',{D}!$B:$B,{dir_ref}' if dir_ref else ''
    return f'SUMIFS({D}!${src}:${src},{D}!$A:$A,">="&{frm},{D}!$A:$A,"<="&{to}{crit})'


def metric_row(ws, r, frm, to, dir_ref):
    """B..I: расход, показы, клики, CTR, CPC, лиды, CPL, CR за период [frm; to]."""
    ws[f'B{r}'] = '=' + sumifs('C', frm, to, dir_ref)
    ws[f'C{r}'] = '=' + sumifs('D', frm, to, dir_ref)
    ws[f'D{r}'] = '=' + sumifs('E', frm, to, dir_ref)
    ws[f'G{r}'] = '=' + sumifs('F', frm, to, dir_ref)
    derived_row(ws, r)


def derived_row(ws, r):
    ws[f'E{r}'] = f'=IFERROR(D{r}/C{r},"")'
    ws[f'F{r}'] = f'=IFERROR(B{r}/D{r},"")'
    ws[f'H{r}'] = f'=IFERROR(B{r}/G{r},"")'
    ws[f'I{r}'] = f'=IFERROR(G{r}/D{r},"")'


# ---------- оформление ----------

def setup_sheet(ws, tab, widths):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[col(i)].width = w


def banner(ws, title_formula, subtitle_formula):
    for r in (1, 2):
        for c in range(1, 40):
            ws.cell(r, c).fill = fill(NAVY)
    ws.row_dimensions[1].height = 34
    ws['A1'] = title_formula
    ws['A1'].font = font(size=18, bold=True, color='FFFFFF')
    ws['A1'].alignment = Alignment(vertical='center')
    ws['A2'] = subtitle_formula
    ws['A2'].font = font(size=10, color='C9D1E3')


def section(ws, r, text):
    ws.cell(r, 1, text).font = font(size=12, bold=True)


def header(ws, r, names):
    for i, n in enumerate(names, 1):
        c = ws.cell(r, i, n)
        c.fill, c.font = fill(HEAD), font(bold=True, color='FFFFFF')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def style_body(ws, r1, r2, ncols, formats, total_row=None):
    for r in range(r1, r2 + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.font = font(bold=(r == total_row))
            cell.fill = fill('E8ECF4' if r == total_row else (BAND if (r - r1) % 2 else 'FFFFFF'))
            cell.border = Border(bottom=thin, top=Side(style='thin', color=NAVY) if r == total_row else None)
            if formats[c - 1]:
                cell.number_format = formats[c - 1]
            cell.alignment = Alignment(horizontal='left' if c == 1 else 'right')


def heat(ws, rng):
    ws.conditional_formatting.add(rng, ColorScaleRule(
        start_type='min', start_color='D8F0E0', mid_type='percentile', mid_value=50,
        mid_color='FFFFFF', end_type='max', end_color='F8D7D5'))


def kpi_block(ws, r, dir_ref):
    """Этот месяц (с 1-го по отчётную дату) против тех же дней прошлого месяца."""
    section(ws, r, 'Этот месяц против тех же дней прошлого')
    h = r + 1
    header(ws, h, ['Период'] + COLS[1:] + ['с', 'по'])
    cur, prev, delta = h + 1, h + 2, h + 3
    ws[f'J{cur}'] = '=DATE(YEAR($E$3),MONTH($E$3),1)'
    ws[f'K{cur}'] = '=$E$3'
    ws[f'J{prev}'] = f'=EDATE(J{cur},-1)'
    ws[f'K{prev}'] = f'=EDATE(K{cur},-1)'
    ws[f'A{cur}'], ws[f'A{prev}'], ws[f'A{delta}'] = 'Этот месяц', 'Прошлый, те же дни', 'Изменение'
    for rr in (cur, prev):
        metric_row(ws, rr, f'$J${rr}', f'$K${rr}', dir_ref)
    for c in range(2, 10):
        L = col(c)
        ws[f'{L}{delta}'] = f'=IFERROR({L}{cur}/{L}{prev}-1,"")'
    style_body(ws, cur, delta, 11, FMT + ['DD.MM.YYYY', 'DD.MM.YYYY'])
    for c in range(1, 12):
        ws.cell(cur, c).font = font(size=14, bold=True)
        ws.cell(prev, c).font = font(color=MUTED)
        ws.cell(delta, c).font = font(bold=True, color=MUTED)
        ws.cell(delta, c).number_format = DELTA_FMT if 1 < c < 10 else 'General'
    ws.row_dimensions[cur].height = 26
    for c in range(2, 10):
        if not GOOD[c - 1]:
            continue
        ref = f'{col(c)}{delta}'
        up, down = ('1A7F37', 'C62828') if GOOD[c - 1] > 0 else ('C62828', '1A7F37')
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'AND(ISNUMBER({ref}),{ref}>0.005)'],
                                                       font=Font(color=up, bold=True)))
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'AND(ISNUMBER({ref}),{ref}<-0.005)'],
                                                       font=Font(color=down, bold=True)))
    return delta + 2


RU_MONTHS = 'CHOOSE(MONTH({x}),"янв","фев","мар","апр","май","июн","июл","авг","сен","окт","ноя","дек")&" "&YEAR({x})'
PERIODS = {
    # (заголовок, число строк, шапка 1-й колонки, начало k-го периода, конец, текстовая подпись)
    'month': ('📅 По месяцам (последние 12)', N_MONTHS, 'Месяц',
              lambda k: f'EDATE(DATE(YEAR($E$3),MONTH($E$3),1),{k - N_MONTHS + 1})',
              lambda a: f'EOMONTH({a},0)', lambda a: RU_MONTHS.format(x=a)),
    'week': ('🗓 По неделям (последние 26)', N_WEEKS, 'Неделя',
             lambda k: f'($E$3-WEEKDAY($E$3,2)+1-7*{N_WEEKS - 1 - k})',
             lambda a: f'({a}+6)', lambda a: f'TEXT({a},"DD.MM")&"–"&TEXT({a}+6,"DD.MM")'),
    'day': ('📆 По дням (последние 60)', N_DAYS, 'День',
            lambda k: f'($E$3-{N_DAYS - 1 - k})', lambda a: a, lambda a: f'TEXT({a},"DD.MM")'),
}


def period_table(ws, r, period, dir_ref, total=False):
    title, n, first_col, start, end, label = PERIODS[period]
    section(ws, r, title)
    header(ws, r + 1, [first_col] + COLS[1:])
    r1 = r + 2
    for k in range(n):
        rr = r1 + k
        a = start(k)
        ws[f'A{rr}'] = '=' + label(a)
        metric_row(ws, rr, a, end(a), dir_ref)
    r2 = r1 + n - 1
    tr = None
    if total:
        tr = r2 + 1
        ws[f'A{tr}'] = 'Итого за 12 мес.'
        for L in 'BCDG':
            ws[f'{L}{tr}'] = f'=SUM({L}{r1}:{L}{r2})'
        derived_row(ws, tr)
    style_body(ws, r1, tr or r2, 9, FMT, tr)
    heat(ws, f'H{r1}:H{r2}')
    return r1, r2, (tr or r2) + 2


# ---------- графики ----------

CHART_COLS = ['M', 'U', 'AC']


def place(ws, chart, grid_row, grid_col, top=4):
    chart.width, chart.height = 13.2, 7.2
    ws.add_chart(chart, f'{CHART_COLS[grid_col]}{top + grid_row * 15}')


def base_chart(kind, title, fmt='#,##0'):
    ch = LineChart() if kind == 'line' else BarChart()
    if kind != 'line':
        ch.type, ch.gapWidth = ('bar' if kind == 'hbar' else 'col'), 40
    ch.title = title
    ch.style = 2
    ch.y_axis.numFmt = fmt
    ch.y_axis.majorGridlines.spPr = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = 'b'
    return ch


def color_series(ch, kind, colors):
    for s, c in zip(ch.series, colors):
        if kind == 'line':
            s.graphicalProperties.line.solidFill = c
            s.graphicalProperties.line.width = 22000
            s.smooth = False
        else:
            s.graphicalProperties.solidFill = c
            s.graphicalProperties.line.solidFill = c


def simple_chart(ws, kind, title, cats_col, val_col, r1, r2, color, cell, fmt='#,##0'):
    ch = base_chart(kind, title, fmt)
    ch.add_data(Reference(ws, min_col=val_col, min_row=r1 - 1, max_row=r2), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cats_col, min_row=r1, max_row=r2))
    color_series(ch, kind, [color])
    ch.legend = None
    place(ws, ch, *cell)


def metric_charts(ws, r1, r2, grid_row, name):
    simple_chart(ws, 'col', f'Расход по {name}, ₽', 1, 2, r1, r2, SPEND, (grid_row, 0))
    simple_chart(ws, 'col', f'Лиды по {name}', 1, 7, r1, r2, LEADS, (grid_row, 1))
    simple_chart(ws, 'line', f'CPL (цена лида) по {name}, ₽', 1, 8, r1, r2, CPL, (grid_row, 2))


def charts_heading(ws, text):
    ws['M3'] = text
    ws['M3'].font = font(size=12, bold=True)


def multi_chart(ws, kind, title, r_head, r1, r2, ncols, cell, stacked=False, fmt='#,##0'):
    """График по направлениям: подписи в A, направления в B.. (без колонки «Итого»)."""
    ch = base_chart(kind, title, fmt)
    ch.add_data(Reference(ws, min_col=2, max_col=1 + ncols, min_row=r_head, max_row=r2), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=r1, max_row=r2))
    if stacked:
        ch.grouping, ch.overlap = 'stacked', 100
    color_series(ch, kind, DIR_COLORS)
    place(ws, ch, *cell)


def by_dir_chart(ws, kind, title, val_col, r1, r2, cell, fmt='#,##0'):
    """Одна величина по направлениям, каждое направление — своим цветом."""
    if kind == 'pie':
        ch = PieChart()
        ch.title = title
        ch.add_data(Reference(ws, min_col=val_col, min_row=r1 - 1, max_row=r2), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=r1, max_row=r2))
        ch.legend.position = 'r'
        from openpyxl.chart.label import DataLabelList
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showPercent = True
        ch.dataLabels.showVal = ch.dataLabels.showCatName = ch.dataLabels.showSerName = False
        ch.dataLabels.showLegendKey = False
    else:
        ch = base_chart(kind, title, fmt)
        ch.add_data(Reference(ws, min_col=val_col, min_row=r1 - 1, max_row=r2), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=r1, max_row=r2))
        ch.legend = None
    s = ch.series[0]
    for i, c in enumerate(DIR_COLORS):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = c
        pt.graphicalProperties.line.solidFill = 'FFFFFF'
        s.dPt.append(pt)
    place(ws, ch, *cell)


# ---------- листы ----------

def top_info(ws, dir_formula):
    ws['A3'], ws['D3'] = ('Направление' if dir_formula else 'Все направления'), 'Данные по'
    if dir_formula:
        ws['B3'] = dir_formula
    ws['E3'] = f'={SET}!$B$7'
    ws['E3'].number_format = 'DD.MM.YYYY'
    for a in ('A3', 'D3'):
        ws[a].font = font(color=MUTED)
    for a in ('B3', 'E3'):
        ws[a].font = font(bold=True)


def direction_sheet(wb, i):
    ws = wb.create_sheet(DIRS[i])
    setup_sheet(ws, DIR_COLORS[i], [18] + [12] * 8 + [11, 11, 3])
    banner(ws, '="🦉 Сова · "&$B$3',
           '="Итоги всех кампаний кабинета VK Реклама · данные по "&TEXT($E$3,"DD.MM.YYYY")')
    top_info(ws, f'={SET}!$A${i + 2}')
    r = kpi_block(ws, 5, '$B$3')
    m1, m2, r = period_table(ws, r, 'month', '$B$3', total=True)
    w1, w2, r = period_table(ws, r, 'week', '$B$3')
    d1, d2, r = period_table(ws, r, 'day', '$B$3')
    charts_heading(ws, '📈 Динамика: по дням · по неделям · по месяцам')
    metric_charts(ws, d1, d2, 0, 'дням')
    metric_charts(ws, w1, w2, 1, 'неделям')
    metric_charts(ws, m1, m2, 2, 'месяцам')
    ws.freeze_panes = 'A4'


def pivot(ws, r, title, period, kind, src=None):
    """Месяц/неделя × направления (+ Итого). kind: spend | leads | cpl (cpl считается из src)."""
    _, n, first_col, start, end, label = PERIODS[period]
    section(ws, r, title)
    h = r + 1
    header(ws, h, [first_col] + [None] * 3 + ['Итого'])
    for j in range(3):
        ws.cell(h, 2 + j).value = f'={SET}!$A${j + 2}'
    r1 = h + 1
    for k in range(n):
        rr = r1 + k
        a = start(k)
        ws[f'A{rr}'] = '=' + label(a)
        for j in range(3):
            L = col(2 + j)
            if kind == 'cpl':
                sp, ld = src
                ws[f'{L}{rr}'] = f'=IFERROR({L}{sp + k}/{L}{ld + k},"")'
            else:
                ws[f'{L}{rr}'] = '=' + sumifs('C' if kind == 'spend' else 'F', a, end(a), f'{L}${h}')
        ws[f'E{rr}'] = (f'=IFERROR(E{src[0] + k}/E{src[1] + k},"")' if kind == 'cpl' else f'=SUM(B{rr}:D{rr})')
    r2 = r1 + n - 1
    f = '#,##0' if kind == 'leads' else '#,##0" ₽"'
    style_body(ws, r1, r2, 5, [None, f, f, f, f])
    for rr in range(r1, r2 + 1):
        ws[f'E{rr}'].font = font(bold=True)
    if kind == 'cpl':
        heat(ws, f'B{r1}:D{r2}')
    return h, r1, r2, r2 + 2


def summary_sheet(wb):
    ws = wb.create_sheet('🦉 Сводная', 0)
    setup_sheet(ws, NAVY, [18] + [12] * 8 + [11, 11, 3])
    banner(ws, '🦉 Сова · Сводная по всем направлениям',
           '="Все кабинеты VK Реклама клиники · данные по "&TEXT($E$3,"DD.MM.YYYY")')
    top_info(ws, None)
    r = kpi_block(ws, 5, None)

    section(ws, r, '🏥 Направления за текущий месяц (те же даты, что выше)')
    header(ws, r + 1, ['Направление'] + COLS[1:] + ['Доля расхода'])
    t1 = r + 2
    for j in range(3):
        rr = t1 + j
        ws[f'A{rr}'] = f'={SET}!$A${j + 2}'
        metric_row(ws, rr, '$J$7', '$K$7', f'$A{rr}')
    tt = t1 + 3
    ws[f'A{tt}'] = 'Итого'
    for L in 'BCDG':
        ws[f'{L}{tt}'] = f'=SUM({L}{t1}:{L}{tt - 1})'
    derived_row(ws, tt)
    for rr in range(t1, tt + 1):
        ws[f'J{rr}'] = f'=IFERROR(B{rr}/$B${tt},"")'
    style_body(ws, t1, tt, 10, FMT + ['0%'], tt)
    by_dir_chart(ws, 'pie', 'Доля расхода, текущий месяц', 2, t1, tt - 1, (0, 0))
    by_dir_chart(ws, 'hbar', 'Лиды по направлениям, текущий месяц', 7, t1, tt - 1, (0, 1))
    by_dir_chart(ws, 'hbar', 'CPL по направлениям, текущий месяц, ₽', 8, t1, tt - 1, (0, 2))
    r = tt + 2

    blocks = {}
    for period, pname, grid in (('month', 'месяцам', 1), ('week', 'неделям', 2)):
        h1, s1, s2, r = pivot(ws, r, f'🧭 Расход по направлениям, по {pname}', period, 'spend')
        h2, l1, l2, r = pivot(ws, r, f'Лиды по направлениям, по {pname}', period, 'leads')
        h3, c1, c2, r = pivot(ws, r, f'CPL по направлениям, по {pname}', period, 'cpl', (s1, l1))
        multi_chart(ws, 'col', f'Расход по {pname}, ₽', h1, s1, s2, 3, (grid, 0), stacked=True)
        multi_chart(ws, 'col', f'Лиды по {pname}', h2, l1, l2, 3, (grid, 1), stacked=True)
        multi_chart(ws, 'line', f'CPL по {pname}, ₽', h3, c1, c2, 3, (grid, 2))

    ws.cell(r, 1, 'Итого по клинике').font = font(size=14, bold=True, color=NAVY)
    d1, d2, r = period_table(ws, r + 1, 'day', None)
    metric_charts(ws, d1, d2, 3, 'дням (вся клиника)')
    charts_heading(ws, '📈 Динамика по направлениям и по клинике в целом')
    ws.freeze_panes = 'A4'


def settings_sheet(wb):
    ws = wb.create_sheet('⚙ Настройки')
    setup_sheet(ws, '9AA3B2', [34, 30, 24, 14])
    header(ws, 1, ['Направление (= название листа)', 'Подключение (заполняет скрипт)', 'Метрика лидов (необяз.)', 'Цвет'])
    for j, n in enumerate(DIRS):
        ws.cell(j + 2, 1, n).font = font(bold=True)
        ws.cell(j + 2, 2, '— (демо-данные)').font = font(color=MUTED)
        ws.cell(j + 2, 4, '#' + DIR_COLORS[j]).fill = fill(DIR_COLORS[j])
        ws.cell(j + 2, 4).font = font(color='FFFFFF', bold=True)
    ws['A7'] = 'Отчёт строится по дату:'
    ws['B7'] = f'=MAX({D}!$A:$A)'
    ws['B7'].number_format = 'DD.MM.YYYY'
    ws['A7'].font, ws['B7'].font = font(bold=True), font(bold=True)
    ws['C7'] = '← последний день в «Данные», считается сам'
    ws['C7'].font = font(color=MUTED)
    notes = [
        'Как пользоваться:',
        '• Впишите названия направлений в A2:A4 (например, «Стоматология») — они подставятся во все листы и графики.',
        '  Вкладки направлений можно переименовать вручную — формулы не сломаются.',
        '• Все цифры берутся с листа «Данные»: 1 строка = 1 день одного направления (итог всех кампаний кабинета).',
        '  Скрипт будет дописывать туда строки каждое утро; пока там демо-данные — удалите строки 2+ перед реальной загрузкой.',
        '• «Этот месяц» = с 1-го числа по отчётную дату; сравнение — с теми же днями прошлого месяца.',
        '• Неделя — с понедельника по воскресенье. CPL = расход / лиды, CR = лиды / клики.',
    ]
    for k, t in enumerate(notes):
        ws.cell(9 + k, 1, t).font = font(bold=(k == 0), color=INK if k == 0 else MUTED)
    ws.freeze_panes = 'A2'


def data_sheet(wb, n_days=365, seed=7):
    ws = wb.create_sheet('Данные')
    setup_sheet(ws, '9AA3B2', [12, 22, 12, 12, 10, 8])
    header(ws, 1, ['Дата', 'Направление', 'Расход', 'Показы', 'Клики', 'Лиды'])
    rnd = random.Random(seed)
    last = date.today() - timedelta(days=1)
    rows = []
    for i, name in enumerate(DIRS):
        budget, cpl = 2500 + 2000 * i, 600 + 350 * i
        for k in range(n_days - 1, -1, -1):
            d = last - timedelta(days=k)
            trend = 1 + 0.5 * (n_days - k) / n_days
            spend = round(budget * trend * (0.7 if d.weekday() >= 5 else 1) * (0.75 + rnd.random() * 0.5))
            shows = round(spend / (0.25 + rnd.random() * 0.1))
            clicks = round(shows * (0.007 + rnd.random() * 0.006))
            leads = max(0, round(spend / (cpl * (0.6 + rnd.random() * 0.8) / trend ** 0.5)))
            rows.append((d, f'={SET}!$A${i + 2}', spend, shows, clicks, leads))
    rows.sort(key=lambda x: (x[0], x[1]))
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font = font()
        ws.cell(r, 1).number_format = 'DD.MM.YYYY'
        ws.cell(r, 3).number_format = '#,##0.00'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:F{len(rows) + 1}'


def main(out):
    wb = Workbook()
    wb.remove(wb.active)
    settings_sheet(wb)
    data_sheet(wb)
    for i in range(3):
        direction_sheet(wb, i)
    summary_sheet(wb)
    order = ['🦉 Сводная'] + DIRS + ['⚙ Настройки', 'Данные']
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.save(out)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'sova-report.xlsx')
